import zipfile
from itertools import zip_longest

from .extractors import unit_id
from .schema import Extraction, Locator, StudyError, Unit


def json_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # Excel date/time caches are real parsed values, not recomputed results.
    return value.isoformat() if hasattr(value, 'isoformat') else str(value)


def extract(path, document, limits):
    import openpyxl
    from defusedxml.ElementTree import iterparse
    from openpyxl.utils import get_column_letter
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    cached = None
    units, warnings, visited = [], [], 0
    try:
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        with zipfile.ZipFile(path) as archive:
            for sheet in formulas.worksheets:
                if sheet.sheet_state != 'visible':
                    warnings.append(f'Hidden sheet excluded: {sheet.title}')
                    continue
                # Bounding the visited rectangle prevents sparse far-away cells causing runaway iteration.
                visited += (sheet.max_row or 0) * (sheet.max_column or 0)
                if visited > limits.max_workbook_cells:
                    raise StudyError('limit', 'Workbook visited cell range exceeds the configured limit.')
                merged = []
                with archive.open(sheet._worksheet_path) as stream:
                    for _, node in iterparse(stream, events=('end',), forbid_dtd=True):
                        if node.tag.endswith('}mergeCell'):
                            merged.append(node.attrib['ref'])
                        node.clear()
                for number, (row, values) in enumerate(zip_longest(sheet.iter_rows(), cached[sheet.title].iter_rows(), fillvalue=()), 1):
                    cells, parts = {}, []
                    for cell, value in zip(row, values):
                        if cell.value is None:
                            continue
                        address = cell.coordinate
                        formula = str(cell.value) if cell.data_type == 'f' else None
                        display = value.value if formula else cell.value
                        if formula and display is None:
                            warnings.append(f'{sheet.title}!{address}: formula has no cached value; not recalculated.')
                        cells[address] = {'formula': formula, 'cached_value': json_value(value.value) if formula else None,
                                          'value': str(cell.value) if cell.value is not None else None}
                        parts.append(f'{address}: {formula} (cached: {display if display is not None else "unavailable"})' if formula else f'{address}: {display}')
                    if not parts:
                        continue
                    cell_range = f'A{number}:{get_column_letter(sheet.max_column)}{number}'
                    units.append(Unit(unit_id(document, f'{sheet.title}:row:{number}'), document.document_id,
                                      len(units), 'row', sheet.title, sheet.title, '\n'.join(parts),
                                      (Locator('xlsx', sheet=sheet.title, cell_range=cell_range, role='row'),),
                                      {'cells': cells, 'merged_ranges': merged}))
    finally:
        formulas.close()
        if cached is not None:
            cached.close()
    warnings.append('Formula caches may be stale; formulas and external links were not executed.')
    return Extraction(units, [], warnings, {'units': len(units), 'text_units': len(units), 'visited_cells': visited}, f'openpyxl-{openpyxl.__version__}')
