import hashlib

from .extractors import unit_id
from .schema import Asset, Extraction, Locator, StudyError, Unit


def _shapes(shapes):
    for shape in shapes:
        if hasattr(shape, 'shapes'):
            yield from _shapes(shape.shapes)
        else:
            yield shape


def extract(path, document, limits):
    import pptx
    from pptx import Presentation
    deck = Presentation(path)
    if len(deck.slides) > limits.max_pages_or_slides:
        raise StudyError('limit', 'Slide count exceeds the configured limit.')
    units, assets, warnings = [], [], []
    for number, slide in enumerate(deck.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title is not None else f'Slide {number}'
        key = unit_id(document, f'slide:{number}')
        text, paragraphs, tables, asset_ids, locations = [], [], [], [], [Locator('pptx', slide=number)]
        for shape in _shapes(slide.shapes):
            locations.append(Locator('pptx', slide=number, shape_id=shape.shape_id))
            coordinates = {name: int(getattr(shape, name)) for name in ('left', 'top', 'width', 'height')}
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    paragraphs.append({'text': paragraph.text, 'level': paragraph.level, 'shape_id': shape.shape_id, 'coordinates_emu': coordinates})
                    if paragraph.text.strip():
                        text.append(('  ' * paragraph.level) + paragraph.text)
            elif shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                tables.append({'shape_id': shape.shape_id, 'rows': rows, 'coordinates_emu': coordinates})
                text.extend(' | '.join(row) for row in rows)
            elif hasattr(shape, 'image'):
                blob, mime = shape.image.blob, shape.image.content_type
                supported = ((mime == 'image/png' and blob.startswith(b'\x89PNG\r\n\x1a\n')) or
                             (mime == 'image/jpeg' and blob.startswith(b'\xff\xd8')) or
                             (mime == 'image/gif' and blob.startswith((b'GIF87a', b'GIF89a'))))
                if supported:
                    coordinates.update(shape_id=shape.shape_id, frame='asset_normalized', slide_width_emu=int(deck.slide_width), slide_height_emu=int(deck.slide_height))
                    asset_id = hashlib.sha256(f'{key}:{shape.shape_id}'.encode() + blob).hexdigest()[:32]
                    asset_ids.append(asset_id)
                    assets.append(Asset(asset_id, key, mime, blob, coordinates))
                warnings.append(f'Slide {number}: embedded image retained where supported; pixels were not interpreted.')
            else:
                warnings.append(f'Slide {number}: unsupported shape/chart {shape.shape_id}; inspect the original deck.')
        body = '\n'.join(text)
        units.append(Unit(key, document.document_id, len(units), 'body', title, '', body,
                          tuple(locations), {'paragraphs': paragraphs, 'tables': tables}, tuple(asset_ids)))
        if not body.strip():
            warnings.append(f'Slide {number}: no usable body text.')
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            notes = frame.text.strip() if frame is not None else ''
            if notes:
                units.append(Unit(unit_id(document, f'notes:{number}'), document.document_id, len(units),
                                  'notes', title, '', notes, (Locator('pptx', slide=number, role='notes'),)))
    return Extraction(units, assets, warnings, {'units': len(units), 'slides': len(deck.slides), 'text_units': sum(bool(u.text.strip()) for u in units)}, f'python-pptx-{pptx.__version__}')
