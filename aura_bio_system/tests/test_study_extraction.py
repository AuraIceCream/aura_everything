import hashlib
import io
import tempfile
import unittest
from pathlib import Path

from aura_bio.study.schema import Limits, StudyError
from aura_bio.study.storage import StudyStore
from aura_bio.study.extractors import extract_document
from study_fixtures import pdf_fixture, xlsx_fixture, pptx_fixture, kerned_pdf_fixture


class StudyExtractionTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.store = StudyStore(self.root / 'packs')
        self.pack = self.store.create_pack('Class')

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def extract(self, name, data):
        document, _ = self.store.add_document(self.pack.pack_id, name, io.BytesIO(data))
        path = self.store.original_path(self.pack.pack_id, document.document_id)
        result = extract_document(path, document, Limits())
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), hashlib.sha256(data).hexdigest())
        return result

    def test_text_preserves_lines_unicode_and_headings(self):
        result = self.extract('notes.md', '# ORION\n\nCalibration is 7 µg.\n'.encode())
        self.assertEqual(result.units[0].locators[0].start_line, 1)
        self.assertEqual(result.units[0].locators[0].end_line, 3)
        self.assertIn('7 µg', result.units[0].text)
        self.assertEqual(result.units[0].title, 'ORION')

    def test_invalid_encoding_fails_clearly(self):
        with self.assertRaises(StudyError) as error:
            self.extract('notes.txt', b'\xff\xfe')
        self.assertEqual(error.exception.code, 'invalid_encoding')

    def test_pdf_keeps_blank_physical_page_number(self):
        path = self.root / 'a.pdf'
        pdf_fixture(path)
        result = self.extract('a.pdf', path.read_bytes())
        self.assertEqual([u.locators[0].page for u in result.units], [1, 2, 3])
        self.assertEqual(result.units[1].text, '')
        self.assertTrue(result.warnings)

    def test_empty_pdf_requires_text_or_ocr(self):
        path = self.root / 'empty.pdf'
        pdf_fixture(path, ('', ''))
        with self.assertRaises(StudyError) as error:
            self.extract('empty.pdf', path.read_bytes())
        self.assertEqual(error.exception.code, 'ocr_required')

    def test_pdf_geometry_recovers_tracked_words_and_preserves_raw_extraction(self):
        import pypdf
        path = self.root / 'kerned.pdf'
        kerned_pdf_fixture(path)
        raw = pypdf.PdfReader(path).pages[0].extract_text().strip()
        self.assertNotIn('Naked mole rats', raw)  # Fixture genuinely reproduces the bug.
        result = self.extract('kerned.pdf', path.read_bytes())
        self.assertIn('Naked mole rats show cancer resistance', result.units[0].text)
        self.assertEqual(result.units[0].metadata['raw_text'], raw)
        self.assertEqual(result.units[0].locators[0].page, 1)

    def test_geometry_does_not_merge_words_in_small_font_prose(self):
        path = self.root / 'small-font.pdf'
        kerned_pdf_fixture(path, font_size=8, letter_gap=0, word_gap=300)
        result = self.extract('small-font.pdf', path.read_bytes())
        self.assertIn('Naked mole rats show cancer resistance', result.units[0].text)

    def test_workbook_formulas_have_no_invented_cache(self):
        path = self.root / 'a.xlsx'
        xlsx_fixture(path)
        result = self.extract('a.xlsx', path.read_bytes())
        formula = result.units[-1].metadata['cells']['A3']
        self.assertEqual(formula['formula'], '=SUM(A1:A2)')
        self.assertIsNone(formula['cached_value'])
        self.assertEqual(result.units[-1].locators[0].cell_range, 'A3:C3')
        self.assertNotIn('Private hidden text', '\n'.join(u.text for u in result.units))
        self.assertIn('B3:C3', result.units[-1].metadata['merged_ranges'])
        self.assertTrue(result.warnings)

    def test_pptx_keeps_notes_table_bullet_level_and_empty_slide(self):
        path = self.root / 'a.pptx'
        pptx_fixture(path)
        result = self.extract('a.pptx', path.read_bytes())
        self.assertEqual({u.locators[0].slide for u in result.units}, {1, 2})
        self.assertTrue(any(u.role == 'notes' and 'Professor note' in u.text for u in result.units))
        bodies = [u for u in result.units if u.role == 'body']
        self.assertIn('ORION', bodies[0].text)
        self.assertTrue(any(p['level'] == 1 for p in bodies[0].metadata['paragraphs']))
        self.assertEqual(bodies[-1].text, '')

    def test_page_limit_fails_instead_of_truncating(self):
        path = self.root / 'a.pdf'
        pdf_fixture(path)
        document, _ = self.store.add_document(self.pack.pack_id, 'a.pdf', io.BytesIO(path.read_bytes()))
        with self.assertRaises(StudyError) as error:
            extract_document(self.store.original_path(self.pack.pack_id, document.document_id), document, Limits(max_pages_or_slides=2))
        self.assertEqual(error.exception.code, 'limit')

    def test_all_blank_deck_is_not_ready_without_usable_text(self):
        from pptx import Presentation
        path = self.root / 'blank.pptx'
        deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(path)
        with self.assertRaises(StudyError) as error:
            self.extract('blank.pptx', path.read_bytes())
        self.assertEqual(error.exception.code, 'no_text')

    def test_grouped_text_and_embedded_image_keep_original_locations(self):
        import base64
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)); box.text = 'Grouped ORION text.'
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jk2kAAAAASUVORK5CYII=')
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(3), width=Inches(1))
        path = self.root / 'image.pptx'; deck.save(path)
        result = self.extract('image.pptx', path.read_bytes())
        self.assertIn('Grouped ORION', result.units[0].text)
        self.assertEqual(result.assets[0].content, png)
        self.assertEqual(result.assets[0].unit_id, result.units[0].unit_id)
        self.assertTrue(result.assets[0].coordinates)

    def test_workbook_cached_date_is_json_safe_and_not_recalculated(self):
        import json
        import zipfile
        from dataclasses import asdict
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active
        sheet['A1'] = '=DATE(2026,1,1)'; sheet['A1'].number_format = 'yyyy-mm-dd'
        original = io.BytesIO(); book.save(original)
        cached = io.BytesIO()
        with zipfile.ZipFile(original) as source, zipfile.ZipFile(cached, 'w') as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    content = content.replace(b'<v></v>', b'<v>46023</v>')
                target.writestr(item, content)
        result = self.extract('cached.xlsx', cached.getvalue())
        encoded = json.dumps(asdict(result.units[0]))
        self.assertIn('2026-01-01', encoded)
