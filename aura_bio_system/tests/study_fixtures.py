"""Tiny source fixtures with literal expectations; never private data."""
from pathlib import Path


def pdf_fixture(path: Path, texts: tuple[str, ...] = ('First lesson.', '', 'Third lesson.')):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        if text:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(f'BT /F1 12 Tf 20 200 Td ({text}) Tj ET'.encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


def kerned_pdf_fixture(path: Path, *, font_size=12, letter_gap=200, word_gap=800):
    """Real glyph coordinates: narrow tracking within words, large word gaps."""
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    words = [f' -{letter_gap} '.join(f'({char})' for char in word) for word in 'Naked mole rats show cancer resistance'.split()]
    stream = DecodedStreamObject()
    stream.set_data((f'BT /F1 {font_size} Tf 20 200 Td [' + f' -{word_gap} '.join(words) + '] TJ ET').encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


def xlsx_fixture(path: Path):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Measurements'
    sheet.append(['Mass', 'g'])
    sheet.append([7, 'g'])
    sheet['A3'] = '=SUM(A1:A2)'
    sheet.merge_cells('B3:C3')
    sheet['B3'] = 'Total'
    hidden = book.create_sheet('Hidden')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = 'Private hidden text'
    book.save(path)


def pptx_fixture(path: Path):
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = 'Lecture one'
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = 'Parent bullet'
    paragraph = box.text_frame.add_paragraph()
    paragraph.text, paragraph.level = 'Nested bullet', 1
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(1, 0).text = 'Term', 'ORION'
    slide.notes_slide.notes_text_frame.text = 'Professor note: fictional course term.'
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
